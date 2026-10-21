"""
Grafico frequenze da Excel - versione 4
---------------------------------------
Legge un file .xlsx con queste colonne:
    A = descrizione
    B = range di frequenza (es. "2400-2483.5 MHz", "5.725-5.875 GHz", "433 MHz")
    C = tipo / categoria (FACOLTATIVA, es. "Wi-Fi", "Video", "Telemetria")
e disegna il grafico delle frequenze con le descrizioni sopra ogni barra.

Novità della v4:
  - campo "Cerca": mostra solo le voci che contengono il testo scritto
    (cerca nella descrizione, nel tipo e nel range, es. "wi-fi" oppure "2.4 ghz");
  - tendina "Tipo": mostra solo le voci di un tipo;
  - "Azzera filtri" rimette tutto. I colori dei tipi restano gli stessi anche
    quando filtri; sovrapposizioni e blocchi sono calcolati sulle voci visibili.

Già presenti nelle versioni precedenti:
  - colori per tipo (colonna C) con legenda;
  - sovrapposizioni evidenziate in rosso e elencate in una tabella;
  - riquadro con i dettagli passando il mouse sopra una barra;
  - grafico spezzato in blocchi quando ci sono grandi vuoti tra i range.

Installazione su Debian/Ubuntu:
    sudo apt install python3-openpyxl python3-matplotlib python3-tk
"""

import re
import textwrap
import tkinter as tk
from dataclasses import dataclass
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from matplotlib import colormaps
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
from matplotlib.figure import Figure
from matplotlib.patches import Patch
from matplotlib.text import Annotation
from matplotlib.ticker import FuncFormatter, MaxNLocator
from openpyxl import load_workbook

# ----------------------------------------------------------------------
# LETTURA DEL FILE
# ----------------------------------------------------------------------

# Tutto viene convertito internamente in MHz
FATTORI = {"hz": 1e-6, "khz": 1e-3, "mhz": 1.0, "ghz": 1e3}
NUMERO_UNITA = re.compile(r"(\d+(?:\.\d+)?)\s*(ghz|mhz|khz|hz)?")
EPS = 1e-6  # tolleranza per i confronti (in MHz)


@dataclass
class Voce:
    descrizione: str
    lo: float  # MHz
    hi: float  # MHz
    riga: int = 0  # numero di riga nel file Excel
    categoria: str = ""  # tipo (colonna C, facoltativa)


def leggi_range(valore, unita_default):
    """Trasforma il testo della cella in (minimo, massimo) in MHz.
    Ritorna None se non riesce a capirlo."""
    if valore is None:
        return None
    if isinstance(valore, (int, float)):
        f = float(valore) * FATTORI[unita_default]
        return f, f

    testo = str(valore).lower().replace(",", ".")
    trovati = NUMERO_UNITA.findall(testo)
    if not trovati:
        return None

    n1, u1 = trovati[0]
    if len(trovati) == 1:  # frequenza singola
        v = float(n1) * FATTORI[u1 or unita_default]
        return v, v

    n2, u2 = trovati[1]
    u_fine = u2 or u1 or unita_default
    u_inizio = u1 or u_fine  # "2.4-2.5 GHz" -> GHz per entrambi
    a = float(n1) * FATTORI[u_inizio]
    b = float(n2) * FATTORI[u_fine]
    return min(a, b), max(a, b)


def leggi_excel(percorso, unita_default):
    """Legge colonna A (descrizione), B (range) e C (tipo, facoltativa) del primo foglio."""
    voci, avvisi = [], []
    wb = load_workbook(percorso, read_only=True, data_only=True)
    ws = wb.active
    for n_riga, riga in enumerate(ws.iter_rows(min_row=1, max_col=3, values_only=True), start=1):
        riga = list(riga) + [None, None, None]
        desc, rng, cat = riga[0], riga[1], riga[2]
        if desc is None and rng is None:
            continue
        parsed = leggi_range(rng, unita_default)
        if parsed is None:
            if n_riga != 1:  # la prima riga di solito è l'intestazione
                avvisi.append(f"riga {n_riga}: range non riconosciuto ({rng!r})")
            continue
        testo = str(desc).strip() if desc is not None else ""
        tipo = str(cat).strip() if cat is not None else ""
        voci.append(Voce(testo or "(senza descrizione)", parsed[0], parsed[1], n_riga, tipo))
    wb.close()
    return voci, avvisi


# ----------------------------------------------------------------------
# BLOCCHI DI FREQUENZA E SOVRAPPOSIZIONI
# ----------------------------------------------------------------------

def soglia_auto(voci):
    """Vuoto massimo automatico: 5% dell'intera estensione dei dati."""
    lo = min(v.lo for v in voci)
    hi = max(v.hi for v in voci)
    return max((hi - lo) * 0.05, 1e-6)


def crea_gruppi(voci, soglia):
    """Raggruppa le voci: se tra una voce e la precedente c'è un vuoto
    maggiore della soglia, comincia un nuovo blocco."""
    ordinate = sorted(voci, key=lambda v: (v.lo, v.hi))
    gruppi = [[ordinate[0]]]
    fine = ordinate[0].hi
    for v in ordinate[1:]:
        if v.lo - fine > soglia:
            gruppi.append([v])
            fine = v.hi
        else:
            gruppi[-1].append(v)
            fine = max(fine, v.hi)
    return gruppi


def trova_sovrapposizioni(voci):
    """Ritorna una lista di (voce_a, voce_b, inizio, fine) per ogni coppia
    di voci i cui range si sovrappongono.
    Due range che si toccano soltanto in un estremo NON si sovrappongono;
    una frequenza singola dentro un range (anche sul bordo) sì."""
    ordinate = sorted(voci, key=lambda v: (v.lo, v.hi))
    coppie = []
    for i, a in enumerate(ordinate):
        for b in ordinate[i + 1:]:
            if b.lo > a.hi + EPS:
                break  # le successive partono ancora più avanti
            inizio = max(a.lo, b.lo)
            fine = min(a.hi, b.hi)
            a_punto = (a.hi - a.lo) <= EPS
            b_punto = (b.hi - b.lo) <= EPS
            if fine - inizio > EPS or (fine - inizio >= -EPS and (a_punto or b_punto)):
                coppie.append((a, b, inizio, max(fine, inizio)))
    return coppie


SENZA_TIPO = "(senza tipo)"

# colori dei tipi: niente rosso, che è riservato alle sovrapposizioni
COLORI_TIPO = ["#1f77b4", "#2ca02c", "#ff7f0e", "#9467bd", "#8c564b", "#e377c2",
               "#17becf", "#bcbd22", "#003f5c", "#00a08a", "#f2b01e", "#6a3d9a"]


def colori_per_tipo(voci):
    """Assegna un colore a ogni tipo (nell'ordine in cui compaiono nel file).
    Le voci senza tipo sono grigie."""
    tipi = []
    for v in voci:
        t = v.categoria or SENZA_TIPO
        if t not in tipi:
            tipi.append(t)
    veri = [t for t in tipi if t != SENZA_TIPO]
    mappa = {t: COLORI_TIPO[i % len(COLORI_TIPO)] for i, t in enumerate(veri)}
    if SENZA_TIPO in tipi:
        mappa[SENZA_TIPO] = "#9e9e9e"
    return mappa


TUTTI_TIPI = "(tutti i tipi)"


def filtra(voci, testo, tipo):
    """Tiene solo le voci che contengono 'testo' (in descrizione, tipo o range)
    e che sono del tipo scelto."""
    testo = (testo or "").strip().lower()
    risultato = []
    for v in voci:
        if tipo and tipo != TUTTI_TIPI and (v.categoria or SENZA_TIPO) != tipo:
            continue
        if testo and not (testo in v.descrizione.lower()
                          or testo in v.categoria.lower()
                          or testo in fmt_range(v).lower()):
            continue
        risultato.append(v)
    return risultato


def fmt_freq(mhz):
    if mhz >= 1000:
        return f"{mhz / 1000:g} GHz"
    if mhz >= 1:
        return f"{mhz:g} MHz"
    return f"{mhz * 1000:g} kHz"


def fmt_range(v):
    if v.hi - v.lo < 1e-9:
        return fmt_freq(v.lo)
    return f"{fmt_freq(v.lo)} – {fmt_freq(v.hi)}"


def fmt_zona(inizio, fine):
    return fmt_range(Voce("", inizio, fine))


# ----------------------------------------------------------------------
# DISEGNO DEL GRAFICO
# ----------------------------------------------------------------------

def testo_tooltip(v, altre):
    """Testo del riquadro che compare passando il mouse sulla barra."""
    righe = [textwrap.fill(v.descrizione, 42), fmt_range(v) + f"   (riga {v.riga})"]
    if v.categoria:
        righe.append(f"Tipo: {v.categoria}")
    if altre:
        righe.append("")
        righe.append("Si sovrappone con:")
        for altra, inizio, fine in altre[:5]:
            righe.append("  • " + textwrap.shorten(altra.descrizione, 34, placeholder="…")
                         + f"  [{fmt_zona(inizio, fine)}]")
        if len(altre) > 5:
            righe.append(f"  ... e altre {len(altre) - 5}")
    return "\n".join(righe)


def disegna(fig, voci, soglia, mostra_range=False, titolo="", per_tipo=True, voci_colori=None):
    """Disegna il grafico su 'fig'.
    Ritorna un dizionario con: n (blocchi), coppie (sovrapposizioni), usa_tipo,
    barre (per il tooltip) e ann (il riquadro del tooltip)."""
    fig.clear()
    gruppi = crea_gruppi(voci, soglia)
    n = len(gruppi)
    coppie = trova_sovrapposizioni(voci)

    # colori: per tipo (se c'è la colonna C) oppure uno diverso per ogni barra
    # (voci_colori = tutte le voci del file: così i colori non cambiano quando si filtra)
    tutte = voci_colori or voci
    usa_tipo = per_tipo and any(v.categoria for v in tutte)
    mappa_tipi = colori_per_tipo(tutte) if usa_tipo else {}

    # per ogni voce: con chi si sovrappone
    altre_di = {}
    for a, b, inizio, fine in coppie:
        altre_di.setdefault(id(a), []).append((b, inizio, fine))
        altre_di.setdefault(id(b), []).append((a, inizio, fine))

    # limiti di ogni blocco (con un po' di margine)
    limiti = []
    for g in gruppi:
        lo = min(v.lo for v in g)
        hi = max(v.hi for v in g)
        span = hi - lo
        pad = span * 0.08 if span > 0 else max(abs(lo) * 0.01, 0.5)
        limiti.append((lo - pad, hi + pad))

    # larghezza dei pannelli proporzionale all'estensione (minimo 16%)
    larghezze = [b - a for a, b in limiti]
    totale = sum(larghezze)
    ratios = [max(w, totale * 0.16) for w in larghezze]
    somma_ratios = sum(ratios)

    righe_max = max(len(g) for g in gruppi)
    fs = 9 if righe_max <= 8 else 8 if righe_max <= 14 else 7
    righe_testo = 2 if righe_max <= 10 else 1
    max_linee = max(1, righe_testo - (1 if mostra_range else 0))

    assi = fig.subplots(1, n, squeeze=False, gridspec_kw={"width_ratios": ratios})[0]
    fig.subplots_adjust(left=0.04, right=0.96, top=0.86, bottom=0.19 if usa_tipo else 0.12, wspace=0.07)
    palette = colormaps["tab10"]
    idx = 0
    barre = []

    marca = dict(marker=[(-1, -0.5), (1, 0.5)], markersize=10, linestyle="none",
                 color="k", mec="k", mew=1, clip_on=False)

    for i, (ax, g, (x0, x1)) in enumerate(zip(assi, gruppi, limiti)):
        ax.set_xlim(x0, x1)
        ax.set_ylim(righe_max - 0.35, -1.1)
        ax.set_yticks([])
        ax.patch.set_visible(False)  # il testo può sconfinare nel blocco vicino
        ax.grid(axis="x", alpha=0.3)
        minw = (x1 - x0) * 0.006

        # zone rosse: sovrapposizioni dentro questo blocco
        ids_gruppo = {id(v) for v in g}
        for a, b, inizio, fine in coppie:
            if id(a) in ids_gruppo:
                ax.axvspan(inizio, max(fine, inizio + minw), color="red", alpha=0.22,
                           linewidth=0, zorder=0.5)

        for r, v in enumerate(g):
            if usa_tipo:
                colore = mappa_tipi[v.categoria or SENZA_TIPO]
            else:
                colore = palette(idx % 10)
            idx += 1
            centro = (v.lo + v.hi) / 2
            larg = max(v.hi - v.lo, minw)
            sovrapposta = id(v) in altre_di
            ax.barh(r, larg, left=centro - larg / 2, height=0.3, color=colore,
                    edgecolor="red" if sovrapposta else "black",
                    linewidth=1.8 if sovrapposta else 0.6, zorder=2)

            barre.append(dict(ax=ax, r=r, x0=centro - larg / 2, x1=centro + larg / 2,
                              tol=(x1 - x0) * 0.012,
                              testo=testo_tooltip(v, altre_di.get(id(v), []))))

            # testo della descrizione sopra la barra
            linee = textwrap.wrap(v.descrizione, width=30) or [""]
            if len(linee) > max_linee:
                linee = linee[:max_linee]
                linee[-1] = linee[-1][:27].rstrip() + "…"
            if mostra_range:
                linee.append(fmt_range(v))

            frac = (centro - x0) / (x1 - x0)
            if frac < 0.33:
                ha, xt = "left", v.lo
            elif frac > 0.67:
                ha, xt = "right", v.hi
            else:
                ha, xt = "center", centro
            ax.text(xt, r - 0.2, "\n".join(linee), ha=ha, va="bottom", fontsize=fs)

        # titolo del blocco = range coperto
        lo = min(v.lo for v in g)
        hi = max(v.hi for v in g)
        ax.set_title(fmt_zona(lo, hi), fontsize=8.5, fontweight="bold")

        # tacche dell'asse
        quota = ratios[i] / somma_ratios
        ax.xaxis.set_major_locator(MaxNLocator(nbins=max(2, min(5, int(quota * 12)))))
        ax.xaxis.set_major_formatter(FuncFormatter(lambda x, _: fmt_freq(x)))
        ax.tick_params(axis="x", labelsize=8)

        # aspetto "asse interrotto"
        ax.spines["top"].set_visible(False)
        if i > 0:
            ax.spines["left"].set_visible(False)
            ax.plot([0, 0], [0, 1], transform=ax.transAxes, **marca)
        if i < n - 1:
            ax.spines["right"].set_visible(False)
            ax.plot([1, 1], [0, 1], transform=ax.transAxes, **marca)

    if titolo:
        fig.suptitle(titolo, fontsize=11)
    if coppie:
        fig.text(0.96, 0.01, "Zone rosse = sovrapposizioni", color="red",
                 ha="right", va="bottom", fontsize=8)

    if usa_tipo:
        presenti = {v.categoria or SENZA_TIPO for v in voci}
        handles = [Patch(facecolor=c, edgecolor="black", label=nome)
                   for nome, c in mappa_tipi.items() if nome in presenti]
        fig.legend(handles=handles, loc="lower left", bbox_to_anchor=(0.04, 0.0),
                   ncol=min(len(handles), 5), fontsize=8, frameon=False)

    # riquadro del tooltip (a livello di figura, così sta sopra a tutto)
    ann = fig.add_artist(Annotation(
        "", xy=(0, 0), xycoords="figure pixels", xytext=(12, 14), textcoords="offset points",
        fontsize=9, ha="left", va="bottom", zorder=50,
        bbox=dict(boxstyle="round,pad=0.4", fc="lightyellow", ec="gray")))
    ann.set_visible(False)

    return dict(n=n, coppie=coppie, barre=barre, ann=ann, usa_tipo=usa_tipo)


def collega_tooltip(canvas, info):
    """Mostra il riquadro quando il mouse passa su una barra.
    Ritorna l'id del collegamento (serve per scollegarlo al ridisegno)."""
    fig = canvas.figure
    ann = info["ann"]

    def su_movimento(ev):
        trovato = None
        if ev.inaxes is not None and ev.xdata is not None and ev.ydata is not None:
            for b in info["barre"]:
                if (b["ax"] is ev.inaxes
                        and b["x0"] - b["tol"] <= ev.xdata <= b["x1"] + b["tol"]
                        and abs(ev.ydata - b["r"]) <= 0.35):
                    trovato = b
                    break

        if trovato is None:
            if ann.get_visible():
                ann.set_visible(False)
                canvas.draw_idle()
            return

        # se siamo vicini ai bordi, il riquadro si apre dall'altra parte
        dx, ha = (12, "left") if ev.x < fig.bbox.width * 0.6 else (-12, "right")
        dy, va = (14, "bottom") if ev.y < fig.bbox.height * 0.7 else (-14, "top")
        ann.set_text(trovato["testo"])
        ann.xy = (ev.x, ev.y)
        ann.xyann = (dx, dy)
        ann.set_ha(ha)
        ann.set_va(va)
        ann.set_visible(True)
        canvas.draw_idle()

    return canvas.mpl_connect("motion_notify_event", su_movimento)


# ----------------------------------------------------------------------
# FINESTRA PRINCIPALE
# ----------------------------------------------------------------------

class App:
    def __init__(self, root):
        self.root = root
        self.percorso = None
        self.voci = []
        self.info = None
        self.cid = None
        self._job = None

        root.title("Grafico frequenze da Excel - v4")
        root.geometry("1280x720")
        root.minsize(900, 560)

        barra = ttk.Frame(root, padding=6)
        barra.pack(fill="x")

        ttk.Button(barra, text="Apri Excel…", command=self.apri).pack(side="left")

        ttk.Label(barra, text="   Unità se manca:").pack(side="left")
        self.var_unita = tk.StringVar(value="MHz")
        cb = ttk.Combobox(barra, textvariable=self.var_unita, width=6, state="readonly",
                          values=["Hz", "kHz", "MHz", "GHz"])
        cb.pack(side="left")
        cb.bind("<<ComboboxSelected>>", lambda e: self.ricarica())

        ttk.Label(barra, text="   Vuoto max per spezzare (MHz, vuoto = auto):").pack(side="left")
        self.var_soglia = tk.StringVar()
        ent = ttk.Entry(barra, textvariable=self.var_soglia, width=8)
        ent.pack(side="left")
        ent.bind("<Return>", lambda e: self.aggiorna())

        self.var_range = tk.BooleanVar(value=False)
        ttk.Checkbutton(barra, text="Mostra range nel testo", variable=self.var_range,
                        command=self.aggiorna).pack(side="left", padx=8)

        self.var_tipo = tk.BooleanVar(value=True)
        ttk.Checkbutton(barra, text="Colora per tipo", variable=self.var_tipo,
                        command=self.aggiorna).pack(side="left", padx=8)

        ttk.Button(barra, text="Aggiorna", command=self.aggiorna).pack(side="left")

        barra2 = ttk.Frame(root, padding=(6, 0, 6, 4))
        barra2.pack(fill="x")
        ttk.Button(barra2, text="Finestra grande", command=self.finestra_grande).pack(side="left")
        ttk.Button(barra2, text="Salva immagine…", command=self.salva).pack(side="left", padx=6)
        self.var_stato = tk.StringVar(value="Apri un file Excel per iniziare.")
        ttk.Label(barra2, textvariable=self.var_stato).pack(side="left", padx=10)

        # filtri: ricerca + tipo
        barra3 = ttk.Frame(root, padding=(6, 0, 6, 4))
        barra3.pack(fill="x")
        ttk.Label(barra3, text="Cerca:").pack(side="left")
        self.var_cerca = tk.StringVar()
        ent_cerca = ttk.Entry(barra3, textvariable=self.var_cerca, width=28)
        ent_cerca.pack(side="left", padx=(4, 12))
        ent_cerca.bind("<KeyRelease>", self.cerca_cambiata)
        ttk.Label(barra3, text="Tipo:").pack(side="left")
        self.var_filtro_tipo = tk.StringVar(value=TUTTI_TIPI)
        self.cb_filtro_tipo = ttk.Combobox(barra3, textvariable=self.var_filtro_tipo, width=22,
                                           state="disabled", values=[TUTTI_TIPI])
        self.cb_filtro_tipo.pack(side="left", padx=(4, 12))
        self.cb_filtro_tipo.bind("<<ComboboxSelected>>", lambda e: self.aggiorna())
        ttk.Button(barra3, text="Azzera filtri", command=self.azzera_filtri).pack(side="left")
        ttk.Label(barra3, foreground="gray",
                  text='   (cerca in descrizione, tipo e range: es. "wi-fi" oppure "2.4 ghz")'
                  ).pack(side="left")

        # tabella delle sovrapposizioni (in basso)
        pannello = ttk.LabelFrame(root, text="Sovrapposizioni", padding=4)
        pannello.pack(side="bottom", fill="x", padx=6, pady=(0, 6))
        self.var_sovr = tk.StringVar(value="Nessun dato")
        ttk.Label(pannello, textvariable=self.var_sovr).pack(anchor="w")
        cornice = ttk.Frame(pannello)
        cornice.pack(fill="x")
        self.tabella = ttk.Treeview(cornice, columns=("a", "b", "zona"), show="headings", height=4)
        self.tabella.heading("a", text="Voce A")
        self.tabella.heading("b", text="Voce B")
        self.tabella.heading("zona", text="Zona in comune")
        self.tabella.column("a", width=420)
        self.tabella.column("b", width=420)
        self.tabella.column("zona", width=260)
        scroll = ttk.Scrollbar(cornice, orient="vertical", command=self.tabella.yview)
        self.tabella.configure(yscrollcommand=scroll.set)
        self.tabella.pack(side="left", fill="x", expand=True)
        scroll.pack(side="right", fill="y")

        self.fig = Figure(figsize=(12, 5.0), dpi=100)
        self.canvas = FigureCanvasTkAgg(self.fig, master=root)
        self.canvas.get_tk_widget().pack(fill="both", expand=True)
        self.disegna_vuoto()

    # --- azioni -------------------------------------------------------
    def scollega_tooltip(self):
        if self.cid is not None:
            self.canvas.mpl_disconnect(self.cid)
            self.cid = None

    def disegna_vuoto(self, messaggio=None):
        self.scollega_tooltip()
        self.info = None
        self.fig.clear()
        self.fig.text(0.5, 0.5, messaggio or "Apri un file Excel\n(colonna A = descrizione, colonna B = range di frequenza)",
                      ha="center", va="center", fontsize=13, color="gray")
        self.canvas.draw_idle()
        self.tabella.delete(*self.tabella.get_children())
        self.var_sovr.set("Nessun dato")

    def apri(self):
        percorso = filedialog.askopenfilename(
            title="Scegli il file Excel",
            filetypes=[("File Excel", "*.xlsx *.xlsm"), ("Tutti i file", "*.*")])
        if not percorso:
            return
        self.percorso = percorso
        self.var_cerca.set("")
        self.var_filtro_tipo.set(TUTTI_TIPI)
        self.ricarica()

    def ricarica(self):
        if not self.percorso:
            return
        try:
            self.voci, avvisi = leggi_excel(self.percorso, self.var_unita.get().lower())
        except Exception as e:
            messagebox.showerror("Errore", f"Non riesco a leggere il file:\n{e}")
            return
        if not self.voci:
            self.disegna_vuoto()
            messagebox.showwarning("Nessun dato",
                                   "Non ho trovato nessun range di frequenza valido nella colonna B.")
            return
        self.aggiorna_filtri_tipo()
        self.aggiorna()
        if avvisi:
            messagebox.showwarning(
                "Alcune righe ignorate",
                "\n".join(avvisi[:8]) + ("\n..." if len(avvisi) > 8 else ""))

    def soglia_corrente(self, voci):
        testo = self.var_soglia.get().strip().replace(",", ".")
        if not testo:
            return soglia_auto(voci)
        try:
            return max(float(testo), 0.0)
        except ValueError:
            messagebox.showerror("Errore", "La soglia deve essere un numero (in MHz).")
            return soglia_auto(voci)

    def aggiorna_filtri_tipo(self):
        """Riempie la tendina dei tipi con quelli presenti nel file."""
        ha_tipi = any(v.categoria for v in self.voci)
        tipi = []
        for v in self.voci:
            t = v.categoria or SENZA_TIPO
            if t not in tipi:
                tipi.append(t)
        valori = [TUTTI_TIPI] + (sorted(tipi, key=str.lower) if ha_tipi else [])
        self.cb_filtro_tipo.configure(values=valori, state="readonly" if ha_tipi else "disabled")
        if self.var_filtro_tipo.get() not in valori:
            self.var_filtro_tipo.set(TUTTI_TIPI)

    def voci_visibili(self):
        return filtra(self.voci, self.var_cerca.get(), self.var_filtro_tipo.get())

    def cerca_cambiata(self, _evento=None):
        """Aggiorna il grafico poco dopo l'ultimo tasto premuto."""
        if self._job is not None:
            self.root.after_cancel(self._job)
        self._job = self.root.after(250, self.aggiorna)

    def azzera_filtri(self):
        self.var_cerca.set("")
        self.var_filtro_tipo.set(TUTTI_TIPI)
        self.aggiorna()

    def aggiorna(self):
        self._job = None
        if not self.voci:
            return
        visibili = self.voci_visibili()
        if not visibili:
            self.disegna_vuoto("Nessuna voce corrisponde ai filtri")
            self.var_stato.set(f"0 di {len(self.voci)} voci: cambia o azzera i filtri")
            return
        soglia = self.soglia_corrente(visibili)
        self.scollega_tooltip()
        self.info = disegna(self.fig, visibili, soglia, self.var_range.get(),
                            Path(self.percorso).name, self.var_tipo.get(), self.voci)
        self.cid = collega_tooltip(self.canvas, self.info)
        self.canvas.draw_idle()
        self.riempi_tabella(self.info["coppie"])
        n_tipi = len({v.categoria for v in visibili if v.categoria})
        extra = f" – {n_tipi} tipi" if n_tipi else ""
        quante = (f"{len(visibili)} voci" if len(visibili) == len(self.voci)
                  else f"{len(visibili)} di {len(self.voci)} voci")
        self.var_stato.set(f"{quante}{extra} – {self.info['n']} blocchi di frequenza "
                           f"(vuoto max: {soglia:g} MHz)")

    def riempi_tabella(self, coppie):
        self.tabella.delete(*self.tabella.get_children())
        for a, b, inizio, fine in sorted(coppie, key=lambda c: c[2]):
            self.tabella.insert("", "end", values=(
                f"{a.descrizione}  [{fmt_range(a)}, riga {a.riga}]",
                f"{b.descrizione}  [{fmt_range(b)}, riga {b.riga}]",
                fmt_zona(inizio, fine)))
        if coppie:
            self.var_sovr.set(f"Trovate {len(coppie)} sovrapposizioni (in rosso nel grafico)")
        else:
            self.var_sovr.set("Nessuna sovrapposizione: i range sono tutti separati")

    def finestra_grande(self):
        if not self.voci:
            messagebox.showinfo("Info", "Apri prima un file Excel.")
            return
        visibili = self.voci_visibili()
        if not visibili:
            messagebox.showinfo("Info", "Nessuna voce corrisponde ai filtri.")
            return
        win = tk.Toplevel(self.root)
        win.title("Grafico frequenze – finestra grande")
        try:
            win.state("zoomed")  # a tutto schermo
        except tk.TclError:
            win.geometry("1300x720")
        fig = Figure(figsize=(14, 7), dpi=100)
        info = disegna(fig, visibili, self.soglia_corrente(visibili), self.var_range.get(),
                       Path(self.percorso).name, self.var_tipo.get(), self.voci)
        canvas = FigureCanvasTkAgg(fig, master=win)
        canvas.draw()
        collega_tooltip(canvas, info)
        toolbar = NavigationToolbar2Tk(canvas, win, pack_toolbar=False)
        toolbar.update()
        toolbar.pack(side="bottom", fill="x")
        canvas.get_tk_widget().pack(side="top", fill="both", expand=True)

    def salva(self):
        if not self.voci or not self.info:
            messagebox.showinfo("Info", "Apri prima un file Excel.")
            return
        percorso = filedialog.asksaveasfilename(
            defaultextension=".png",
            filetypes=[("PNG", "*.png"), ("PDF", "*.pdf"), ("SVG", "*.svg")])
        if percorso:
            self.info["ann"].set_visible(False)  # niente tooltip nell'immagine
            self.fig.savefig(percorso, dpi=200)
            self.var_stato.set(f"Salvato: {percorso}")


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
